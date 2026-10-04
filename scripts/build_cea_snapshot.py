#!/usr/bin/env python3
"""Rebuild a date-ordered CEA All India series from archived daily XLSX reports.

Requires: python -m pip install -r scripts/requirements.txt
The app itself is Node-only and reads the committed JSON snapshot.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "references" / "cea_reports"
OUTPUT = ROOT / "lib" / "cea_snapshot.json"
REPORT_PAGE = "https://gen-re.cea.gov.in/reports"
REPORT_URL = "https://gen-re.cea.gov.in/public/uploads/dailyReport/excel/Report-{date}.xlsx"
SHEETS = "Generation"
ROW = 47
VALUE_COLUMNS = {"Wind": "B", "Solar": "C", "Other RES": "D", "Reported total": "E"}
UNIT = "MU net"


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalize_for_hash(value: Any) -> Any:
    """Use language-neutral JSON numbers (e.g. 526, not 526.0)."""
    if isinstance(value, list):
        return [normalize_for_hash(item) for item in value]
    if isinstance(value, dict):
        return {key: normalize_for_hash(item) for key, item in value.items()}
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def file_date(path: Path) -> date | None:
    match = re.fullmatch(r"Report-(\d{4}-\d{2}-\d{2})\.xlsx", path.name)
    return date.fromisoformat(match.group(1)) if match else None


def build(start: date | None = None, end: date | None = None) -> dict[str, Any]:
    files = [(file_date(path), path) for path in REPORTS.glob("Report-*.xlsx")]
    files = [(period, path) for period, path in files if period is not None]
    if start:
        files = [(period, path) for period, path in files if period >= start]
    if end:
        files = [(period, path) for period, path in files if period <= end]
    files.sort(key=lambda item: item[0])
    if not files:
        raise SystemExit("No daily CEA Excel reports found in references/cea_reports.")
    if start and files[0][0] != start:
        raise SystemExit(f"Requested start date {start} is missing from the archived reports.")
    if end and files[-1][0] != end:
        raise SystemExit(f"Requested end date {end} is missing from the archived reports.")
    for index, (period, _) in enumerate(files):
        expected = files[0][0] + timedelta(days=index)
        if period != expected:
            raise SystemExit(f"Archive has a missing or out-of-order report date near {expected}; found {period}.")

    snapshot_built_at = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    observations: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    reconciliation_differences: list[float] = []
    used_dates: set[str] = set()

    for period, path in files:
        period_iso = period.isoformat()
        if period_iso in used_dates:
            raise SystemExit(f"Duplicate report date: {period_iso}")
        used_dates.add(period_iso)
        workbook = load_workbook(path, data_only=True, read_only=True)
        if SHEETS not in workbook.sheetnames:
            raise SystemExit(f"{path.name}: expected worksheet {SHEETS!r} not found")
        sheet = workbook[SHEETS]
        label = str(sheet["A47"].value or "")
        if "All India" not in label:
            raise SystemExit(f"{path.name}: A47 is not the All India total row: {label!r}")
        date_header = str(sheet["B5"].value or "").strip()
        try:
            header_date = datetime.strptime(date_header, "%d-%B-%y").date()
        except ValueError as error:
            raise SystemExit(f"{path.name}: unrecognized report date header {date_header!r}") from error
        if header_date != period:
            raise SystemExit(f"{path.name}: filename date {period} does not match sheet header {header_date}")

        values: dict[str, float] = {}
        for label_name, column in VALUE_COLUMNS.items():
            raw = sheet[f"{column}{ROW}"].value
            if not isinstance(raw, (int, float)):
                raise SystemExit(f"{path.name}: {column}{ROW} is not a numeric value: {raw!r}")
            values[label_name] = float(raw)
        component_sum = sum(values[name] for name in ("Wind", "Solar", "Other RES"))
        reconciliation_difference = round(component_sum - values["Reported total"], 5)
        if abs(reconciliation_difference) > 0.02:
            raise SystemExit(f"{path.name}: source total does not reconcile; difference={reconciliation_difference}")
        reconciliation_differences.append(abs(reconciliation_difference))

        raw_bytes = path.read_bytes()
        report_id = f"cea-daily-{period_iso}"
        source_url = REPORT_URL.format(date=period_iso)
        component = {
            "report_id": report_id,
            "report_name": f"DailyRE{period.strftime('%d%m%Y')}",
            "report_date": period_iso,
            "source_url": source_url,
            "file_name": path.name,
            "local_path": f"references/cea_reports/{path.name}",
            "sha256": sha256(raw_bytes),
            "size_bytes": len(raw_bytes),
            "snapshot_built_at": snapshot_built_at,
            "sheet": SHEETS,
            "all_india_row": ROW,
            "row_label": label,
            "report_date_header": date_header,
            "unit_label": UNIT,
            "cells": {name: f"{column}{ROW}" for name, column in VALUE_COLUMNS.items()},
            "reported_total": round(values["Reported total"], 5),
            "extracted_components_total": round(component_sum, 5),
            "reconciliation_difference": reconciliation_difference,
            "reconciliation_status": "passed",
        }
        components.append(component)
        for technology in ("Wind", "Solar", "Other RES"):
            column = VALUE_COLUMNS[technology]
            observations.append({
                "observation_id": f"{report_id}-{technology.lower().replace(' ', '-')}",
                "source_id": "src-cea-90d",
                "date": period_iso,
                "technology": technology,
                "geography": "All India",
                "value": round(values[technology], 5),
                "unit": "MU",
                "quality_status": "verified_report_extraction",
                "source_report_id": report_id,
                "locator": {
                    "type": "xlsx_cell",
                    "report_id": report_id,
                    "sheet": SHEETS,
                    "row": ROW,
                    "cell": f"{column}{ROW}",
                    "report_date_header": date_header,
                    "reported_total_cell": "E47",
                },
            })
        workbook.close()

    canonical = json.dumps(normalize_for_hash(observations), ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    manifest = {
        "schema_version": 1,
        "dataset_id": "src-cea-90d",
        "dataset_version": f"dv-cea-{files[0][0].isoformat()}-{files[-1][0].isoformat()}-v1",
        "title": "CEA Daily Renewable Generation · All India · 90-report snapshot",
        "publisher": "Central Electricity Authority, Ministry of Power, Government of India",
        "source_url": REPORT_PAGE,
        "coverage_start": files[0][0].isoformat(),
        "coverage_end": files[-1][0].isoformat(),
        "report_count": len(components),
        "observation_count": len(observations),
        "technology_labels": ["Wind", "Solar", "Other RES"],
        "unit": UNIT,
        "method": "Read only the All India row (47), daily columns B–D, from each archived daily report's Generation sheet; retain source row/cell and report-level SHA-256. Column E is used only for source-total reconciliation.",
        "snapshot_built_at": snapshot_built_at,
        "components": components,
        "reconciliation": {
            "definition": "Wind + Solar + Other RES compared with CEA's published All India daily total, per report.",
            "reports_reconciled": sum(1 for component in components if component["reconciliation_status"] == "passed"),
            "reports_total": len(components),
            "tolerance_mu": 0.02,
            "max_abs_difference_mu": round(max(reconciliation_differences, default=0), 5),
            "status": "passed",
        },
        "normalized_observations_sha256": sha256(canonical),
        "caveat": "Historical snapshot of published daily reports; not live data. RE-PREDICT is independent and not endorsed by CEA. Other RES is CEA's combined category; do not disaggregate it into biomass, bagasse, and small hydro without underlying data.",
    }
    OUTPUT.write_text(json.dumps({"manifest": manifest, "observations": observations}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--start", type=date.fromisoformat)
    parser.add_argument("--end", type=date.fromisoformat)
    args = parser.parse_args()
    result = build(args.start, args.end)
    print(json.dumps({
        "output": str(OUTPUT),
        "coverage": [result["coverage_start"], result["coverage_end"]],
        "reports": result["report_count"],
        "observations": result["observation_count"],
        "reconciliation": result["reconciliation"],
        "normalized_observations_sha256": result["normalized_observations_sha256"],
    }, indent=2))
